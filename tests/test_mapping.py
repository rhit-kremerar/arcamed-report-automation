import csv
import importlib.util
import tempfile
import sys
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('mapping', ROOT / 'src/mapping.py')
mapping = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mapping)
sys.path.insert(0, str(ROOT / 'src'))
import app


class ConversionTests(unittest.TestCase):
    def test_august_ui_writes_dr_and_preserves_other_months(self):
        """Use the supplied exports to test August through the complete UI workflow."""
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source = folder / 'financials.xlsx'
            template = folder / 'RE Only.xlsx'
            for csv_path, output, tab in [
                (ROOT / 'data/50001 Arcamed 08.csv', source, 'Report1 (6)'),
                (ROOT / 'data/FINANCIALS 2026-08-31 - RE Only(TB RE) (1).csv', template, 'TB RE')]:
                book = Workbook()
                sheet = book.active
                sheet.title = tab
                for row in mapping.read_rows(csv_path, ''):
                    sheet.append(row)
                if tab == 'TB RE':
                    sheet['DR6'] = -123  # Make the August replacement observable.
                    sheet['DQ6'] = 456  # July must stay untouched.
                    book.create_sheet('MASTER DEFINITION')['B1'] = 7
                book.save(output)
                book.close()
            output, audit, report = app.convert(dict(
                input=str(source), template=str(template),
                mapping=str(ROOT / 'data/50001_to_TB_RE_Mapping.csv'), folder='',
                source_sheet='Report1 (6)', mapping_sheet='Mappings',
                blank_policy='Leave empty (default)', allow_unmapped=True))
            self.assertEqual(report['period'], '2026-08')
            self.assertEqual(report['month_column'], 122)
            self.assertTrue(report['output_written'])
            book = load_workbook(output)
            self.assertEqual(book['TB RE']['DR6'].value, 88418.99)
            self.assertEqual(book.active.title, 'TB RE')
            self.assertTrue(book.views[0].showSheetTabs)
            self.assertEqual(book['TB RE'].sheet_view.selection[0].activeCell, 'DR4')
            self.assertEqual(book['TB RE']['DQ6'].value, 456)
            self.assertEqual(book['MASTER DEFINITION']['B1'].value, 8)
            book.close()
            book = load_workbook(template)
            self.assertEqual(book['TB RE']['DR6'].value, -123)
            book.close()

    def test_ui_filename_and_blank_default(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            book = load_workbook(template)
            book['TB RE'].append(['PL', '8888', 'Missing numeric value', 6, 75])
            book.save(template)
            book.close()
            options = dict(input=str(source), template=str(template), mapping=str(mapping_file),
                           folder='', source_sheet='Report1 (6)', mapping_sheet='Mappings',
                           blank_policy='Leave empty (default)', allow_unmapped=False)
            output, audit, report = app.convert(options)
            self.assertEqual(output.name, 'TB_September.xlsx')
            self.assertEqual(output.parent, source.parent)
            self.assertTrue(report['output_written'])
            self.assertIsNone(audit)
            self.assertEqual(list(folder.glob('*_audit.json')), [])
            book = load_workbook(output)
            self.assertIsNone(book['TB RE']['E7'].value)
            self.assertEqual(book['TB RE']['D7'].value, 6)
            self.assertEqual(book['TB RE']['E6'].value, '=SUM(E4:E5)')
            book.close()
            second, _, _ = app.convert(options)
            self.assertEqual(second.name, 'TB_September_1.xlsx')

    def test_ui_audit_mode_saves_only_numbered_audits(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            options = dict(input=str(source), template=str(template), mapping=str(mapping_file),
                           folder='', source_sheet='Report1 (6)', mapping_sheet='Mappings',
                           blank_policy='Leave empty (default)', allow_unmapped=False, mode='json')
            output, audit, report = app.convert(options)
            self.assertFalse(output.exists())
            self.assertTrue(audit.exists())
            self.assertFalse(report['output_written'])
            _, second_audit, _ = app.convert(options)
            self.assertEqual(second_audit.name, 'TB_September_1_audit.json')
            options['mode'] = 'pdf'
            _, pdf, pdf_report = app.convert(options)
            self.assertEqual(pdf.suffix, '.pdf')
            self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
            self.assertFalse(pdf_report['output_written'])
            options['mode'] = 'workbook'
            output, audit, _ = app.convert(options)
            self.assertEqual(output.name, 'TB_September.xlsx')
            self.assertIsNone(audit)

    def test_ui_blocked_normal_run_saves_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            options = dict(input=str(source), template=str(template), mapping=str(mapping_file),
                           folder='', source_sheet='Report1 (6)', mapping_sheet='Mappings',
                           blank_policy='Stop if amounts are missing', allow_unmapped=False)
            output, audit, report = app.convert(options)
            self.assertTrue(report['write_blockers'])
            self.assertFalse(output.exists())
            self.assertIsNone(audit)
            self.assertEqual(list(folder.glob('*_audit.json')), [])

    def excel_mapping(self, csv_path, output, sheet_name='Mappings'):
        """Model a client saving the same mapping rows in an Excel workbook."""
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        with csv_path.open(newline='', encoding='utf-8-sig') as handle:
            for row in csv.reader(handle):
                sheet.append(row)
        book.save(output)
        book.close()
        return output

    def test_excel_mapping_matches_csv_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            csv_path = ROOT / 'data/50001_to_TB_RE_Mapping.csv'
            excel_path = self.excel_mapping(csv_path, Path(tmp) / 'mapping.xlsx')
            source = mapping.read_rows(ROOT / 'data/50001 Arcamed 08.csv', '')
            target = mapping.read_rows(ROOT / 'data/FINANCIALS 2026-08-31 - RE Only(TB RE) (1).csv', '')
            self.assertEqual(mapping.build_report(source, target, csv_path),
                             mapping.build_report(source, target, excel_path))

    def test_client_appended_excel_row_is_used_in_conversion(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, csv_path = self.fixture(folder)
            excel_path = self.excel_mapping(csv_path, folder / 'mapping.xlsx', 'Client rules')
            book = load_workbook(excel_path)
            sheet = book['Client rules']
            sheet.delete_rows(3)
            sheet['G1'] = 'Client notes'
            sheet['G4'] = 'Empty mapping row with a note is ignored'
            # Numeric source IDs and numeric destination row hints are normal in Excel.
            sheet.append([400007, 'Rent', '4500', 'Rent', 5, 'debit minus credit', 'Added by client'])
            book.save(excel_path)
            book.close()
            output = folder / 'result.xlsx'
            result = mapping.main(['--input', str(source), '--template', str(template),
                                   '--mapping', str(excel_path), '--mapping-sheet', 'Client rules',
                                   '--output', str(output), '--report', str(folder / 'report.json'),
                                   '--missing-lines', 'preserve'])
            self.assertEqual(result, 0)
            book = load_workbook(output)
            self.assertEqual(book['TB RE']['E5'].value, -10)
            book.close()

    def test_excel_mapping_rejects_duplicates_and_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            _, _, csv_path = self.fixture(folder)
            excel_path = self.excel_mapping(csv_path, folder / 'mapping.xlsx')
            book = load_workbook(excel_path)
            sheet = book['Mappings']
            sheet.append([100106, 'Cash', '1110', 'Cash', 4, 'ending balance'])
            sheet.append([100405, 'Receivables', '=1+1', 'Receivables', 6, 'ending balance'])
            book.save(excel_path)
            book.close()
            errors = []
            mapping.read_mapping(excel_path, errors)
            self.assertTrue(any('Duplicate mapping' in error for error in errors))
            self.assertTrue(any('not formulas' in error for error in errors))
            with self.assertRaisesRegex(ValueError, 'needs a'):
                mapping.read_mapping(excel_path, [], 'Missing tab')

    def test_august_exports(self):
        report, values, column = mapping.build_report(
            mapping.read_rows(ROOT / 'data/50001 Arcamed 08.csv', ''),
            mapping.read_rows(ROOT / 'data/FINANCIALS 2026-08-31 - RE Only(TB RE) (1).csv', ''),
            ROOT / 'data/50001_to_TB_RE_Mapping.csv')
        self.assertEqual(report['errors'], [])
        self.assertEqual(report['source_count'], 77)
        self.assertEqual(len(report['unmapped_accounts']), 13)
        self.assertEqual(str(report['totals']['debit']), '2319336.99')
        self.assertEqual(str(values[6]), '88418.99')
        self.assertEqual(column, 122)  # DR; hidden YTD August header is excluded.
        differences = [v for v in report['comparisons'] if v['difference'] not in (None, 0)]
        self.assertEqual([v['row'] for v in differences], [26])

    def fixture(self, folder):
        source = folder / 'source.xlsx'
        template = folder / 'template.xlsx'
        mapping_file = folder / 'mapping.csv'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Report1 (6)'
        sheet['A3'] = 'Period = Sep 2026'
        for row, values in enumerate([
                [100106, 'Cash', 0, 10, 0, 10],
                [400007, 'Rent', 0, 0, 10, -10],
                ['Total', '', 0, 10, 10, 0]], 7):
            for column, value in enumerate(values, 1):
                sheet.cell(row, column, value)
        book.save(source)
        book = Workbook()
        sheet = book.active
        sheet.title = 'TB RE'
        sheet.append([])
        sheet.append([])
        sheet.append(['Cat', 'Acc', 'Desc', "AUG '26", "SEP '26"])
        sheet.append(['BSA', '1110', 'Cash', 42, '=1+1'])
        sheet.append(['PL', '4500', 'Rent', -42, 99])
        sheet.append(['PL', '9999', 'Absent', 5, '=SUM(E4:E5)'])
        sheet['E4'].number_format = '#,##0.00'
        book.create_sheet('MASTER DEFINITION')['B1'] = 8
        book.create_sheet('PL-RE')['A1'] = "='TB RE'!E5"
        book.save(template)
        with mapping_file.open('w', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(['50001 account', '50001 description', 'TB RE account',
                             'TB RE description', 'TB RE row', 'Amount to use'])
            writer.writerow(['100106', 'Cash', '1110', 'Cash', 4, 'ending balance'])
            writer.writerow(['400007', 'Rent', '4500', 'Rent', 5, 'debit minus credit'])
        return source, template, mapping_file

    def test_copy_preserves_history_formulas_and_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            output = folder / 'result.xlsx'
            result = mapping.main(['--input', str(source), '--template', str(template),
                                   '--mapping', str(mapping_file), '--output', str(output),
                                   '--report', str(folder / 'report.json'), '--missing-lines', 'preserve'])
            self.assertEqual(result, 0)
            book = load_workbook(output)
            self.assertEqual(book['TB RE']['D4'].value, 42)
            self.assertEqual(book['TB RE']['E4'].value, 10)
            self.assertEqual(book['TB RE']['E5'].value, -10)
            self.assertEqual(book['TB RE']['E6'].value, '=SUM(E4:E5)')
            self.assertEqual(book['TB RE']['E4'].number_format, '#,##0.00')
            self.assertEqual(book['PL-RE']['A1'].value, "='TB RE'!E5")
            self.assertEqual(book['MASTER DEFINITION']['B1'].value, 9)
            self.assertTrue(book.calculation.fullCalcOnLoad)
            book.close()
            original = load_workbook(template)
            self.assertEqual(original['TB RE']['E4'].value, '=1+1')
            original.close()

    def test_default_missing_lines_blocks_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            output = folder / 'result.xlsx'
            self.assertEqual(mapping.main(['--input', str(source), '--template', str(template),
                                          '--mapping', str(mapping_file), '--output', str(output),
                                          '--report', str(folder / 'report.json')]), 1)
            self.assertFalse(output.exists())

    def test_moved_row_and_duplicate_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            source, template, mapping_file = self.fixture(folder)
            rows = mapping.read_rows(template, 'TB RE')
            rows[3], rows[4] = rows[4], rows[3]
            source_rows = mapping.read_rows(source, 'Report1 (6)')
            report, values, _ = mapping.build_report(source_rows, rows, mapping_file)
            self.assertEqual(len(report['moved_destinations']), 2)
            self.assertEqual(values[5], 10)
            source_rows.insert(7, source_rows[6])
            report, _, _ = mapping.build_report(source_rows, rows, mapping_file)
            self.assertTrue(any('Duplicate source' in error for error in report['errors']))


if __name__ == '__main__':
    unittest.main()
